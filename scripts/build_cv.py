# Genera el CV de Mario Huarte Nolasco en .docx con python-docx, en inglés o español.
# Uso: python3 build_cv.py [en|es]  → Mario_Huarte_CV.docx | Mario_Huarte_CV_ES.docx
# Contenido: datos confirmados por Mario (sept 2026). Una columna, tamaño carta.
# El mismo contenido vive en cv.html / cv_es.html (fuente del PDF); mantener ambos alineados.
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LANG = (sys.argv[1] if len(sys.argv) > 1 else 'en').lower()
OUT = 'Mario_Huarte_CV.docx' if LANG == 'en' else 'Mario_Huarte_CV_ES.docx'

# ── Textos por idioma ─────────────────────────────────────────────
# Cada bullet es una lista de (texto, negrita)
T = {
 'en': {
  'role': 'MDM Specialist | TIBCO EBX Developer | Java | Data Governance',
  'contact1': 'Tel: 55 6229 3691     Email: mario_huarte@outlook.com     LinkedIn: linkedin.com/in/mario-huarte-nolasco-874436202',
  'contact2': 'Mexico City, Mexico     Portfolio: mariounam.github.io     Cédula Profesional: 14265872',
  'summary': [('Computer Engineer (UNAM) with ',False),('5 years of experience',True),(' in Master Data Management (MDM) and ',False),('TIBCO EBX',True),
   (' implementation, since December 2021. Five MDM implementations for clients in retail, manufacturing, financial services, pharma and telecom, modeling customer, employee and location domains on datasets of up to ',False),('200M records',True),
   ('. Specialized in master data modeling, ',False),('Java',True),(' extensions on EBX, Data Governance workflows, Data Quality and Match & Merge rules. Daily use of ',False),('SQL',True),
   ('; integration through ',False),('SnapLogic iPaaS',True),(', Oracle, PostgreSQL, SQL Server and AWS. Occasional use of ',False),('Python',True),(' for process automation. English B2 — technical proficiency.',False)],
  'h_comp': 'Key Competencies',
  'comp': 'Master Data Management (MDM)  |  TIBCO EBX 5.x–6.x  |  EBX Add-ons & Custom Modules  |  Java  |  Data Governance  |  Data Quality  |  Data Cleansing  |  Match & Merge  |  Data Stewardship  |  Reference Data Management  |  Data Modeling  |  EBX Workflows  |  EBX Data Spaces & Datasets  |  SnapLogic iPaaS  |  REST Services  |  SQL  |  Oracle  |  PostgreSQL  |  SQL Server  |  AWS  |  Azure (AZ-900)  |  Python  |  Power BI  |  Enterprise Data Management',
  'h_exp': 'Professional Experience',
  'r1': ('MDM Consultant | TIBCO EBX Developer','Alldatum Business  ·  Mexico City, Mexico','Aug 2023 – Present','3 years 2 months'),
  'r1b': [
   [('Implementation and configuration of MDM solutions on ',False),('TIBCO EBX',True),(': data model design, Data Spaces and Datasets for customer, employee and location domains, on datasets of up to 200M records.',False)],
   [('Data cleansing and quality rules in ',False),('TIBCO EBX',True),(': validations, constraints and corrections on master data records within the platform.',False)],
   [('Java extensions on EBX: triggers, constraints, custom validations, REST services, workflow scripts and importers — created from scratch or evolved from existing modules.',False)],
   [('EBX integration with external systems through ',False),('SnapLogic iPaaS',True),(' and native EBX REST services: Oracle, PostgreSQL, SQL Server and AWS.',False)],
   [('Telecom: designed cleansing and ',False),('Match & Merge',True),(" rules for customer data from two Central American countries, after researching each country's phone numbering plans and address standards; delivered clean, deduplicated data back to the operator.",False)],
   [('Data quality monitoring and reporting with ',False),('Power BI',True),(' and ',False),('SQL',True),(' to track governance KPIs for client projects.',False)],
   [('Five MDM implementations across ',False),('retail, manufacturing, financial services, pharma and telecom',True),(', working with business and IT teams to translate requirements into EBX technical configurations.',False)]],
  'r2': ('Intern – Data and MDM Consultant','Alldatum Business  ·  Mexico City, Mexico','Dec 2021 – Aug 2023','1 year 9 months'),
  'r2b': [[('Initial learning and configuration of the TIBCO EBX platform under senior consultant supervision.',False)],
          [('Support in identifying data sources and mapping them to MDM models during implementation projects.',False)]],
  'r3': ('Programming Instructor','PROTECO  ·  School of Engineering, UNAM','Feb 2021 – Feb 2022','1 year 1 month'),
  'r3b': [[('Taught 10+ online courses and workshops (OOP, web development, Python, Linux) to groups of 20–30 Computer Engineering students, including onboarding of new program members.',False)],
          [('Created learning materials and one-on-one tutoring focused on technical problem-solving, plus practical workshops such as building a professional CV.',False)]],
  'h_edu': 'Education',
  'e1': ('Bachelor of Engineering in Computer Science (Ingeniero en Computación)','Facultad de Ingeniería · Universidad Nacional Autónoma de México (UNAM)','2017 – 2024','Degree & license 2024'),
  'e2': ('Software Development Technician (Técnico en Desarrollo de Software)','Colegio de Ciencias y Humanidades, Azcapotzalco · UNAM','Aug 2014 – Mar 2017','Graduated'),
  'h_cert': 'Certifications',
  'certs': [[('Google AI',True),(' — Google · Aug 2026 · Credential ID LY3LGZVG593W',False)],
            [('Diploma: Desarrollo de Habilidades Directivas (240 h)',True),(' — Facultad de Ingeniería, UNAM · Mar 2024 · Folio 4283',False)],
            [('SnapLogic Integrator Certification',True),(' — SnapLogic · Aug 2023 · verify.skilljar.com/c/ewn4n2hoqxh5',False)],
            [('Google Data Analytics Professional Certificate',True),(' — Google / Coursera · Mar 2023 · coursera.org/share/ef3d02dd5eafc5509d940ecdb1be61b2',False)],
            [('Microsoft Certified: Azure Fundamentals (AZ-900)',True),(' — Microsoft via Certiport · Jun 2020',False)]],
  'udemy': [('Continuous training (Udemy, 2025): ',True),('Curso Práctico SQL Server: De Principiante a Experto · SQL Consultas en SQL Server · Administrador Linux – Curso completo desde cero (LPIC-1) · Ciberseguridad – CompTIA Security+ Módulo 1 · Ciberseguridad: Protege tu información de cibercriminales. Official TIBCO EBX platform training.',False)],
  'h_lang': 'Languages',
  'langs': 'Spanish — Native     |     English — B2 (technical proficiency)',
 },
 'es': {
  'role': 'Especialista MDM | Desarrollador TIBCO EBX | Java | Gobierno de Datos',
  'contact1': 'Tel: 55 6229 3691     Email: mario_huarte@outlook.com     LinkedIn: linkedin.com/in/mario-huarte-nolasco-874436202',
  'contact2': 'Ciudad de México, México     Portafolio: mariounam.github.io     Cédula Profesional: 14265872',
  'summary': [('Ingeniero en Computación (UNAM) con ',False),('5 años de experiencia',True),(' en Master Data Management (MDM) e implementación de ',False),('TIBCO EBX',True),
   (', desde diciembre de 2021. Cinco implementaciones MDM para clientes de retail, manufactura, servicios financieros, farmacéutico y telecomunicaciones, modelando dominios de cliente, empleado y ubicaciones sobre conjuntos de hasta ',False),('200M de registros',True),
   ('. Especializado en modelado de datos maestros, extensiones ',False),('Java',True),(' sobre EBX, workflows de Gobierno de Datos, reglas de Calidad de Datos y Match & Merge. Uso diario de ',False),('SQL',True),
   ('; integración mediante ',False),('SnapLogic iPaaS',True),(', Oracle, PostgreSQL, SQL Server y AWS. Uso ocasional de ',False),('Python',True),(' para automatización de procesos. Inglés B2 — competencia técnica.',False)],
  'h_comp': 'Competencias clave',
  'comp': 'Master Data Management (MDM)  |  TIBCO EBX 5.x–6.x  |  Add-ons EBX y módulos personalizados  |  Java  |  Gobierno de datos  |  Calidad de datos  |  Limpieza de datos  |  Match & Merge  |  Data Stewardship  |  Gestión de datos de referencia  |  Modelado de datos  |  Workflows EBX  |  Data Spaces y Datasets EBX  |  SnapLogic iPaaS  |  Servicios REST  |  SQL  |  Oracle  |  PostgreSQL  |  SQL Server  |  AWS  |  Azure (AZ-900)  |  Python  |  Power BI  |  Enterprise Data Management',
  'h_exp': 'Experiencia profesional',
  'r1': ('Consultor MDM | Desarrollador TIBCO EBX','Alldatum Business  ·  Ciudad de México','Ago 2023 – Presente','3 años 2 meses'),
  'r1b': [
   [('Implementación y configuración de soluciones MDM en ',False),('TIBCO EBX',True),(': diseño del modelo de datos, Data Spaces y Datasets para dominios de cliente, empleado y ubicaciones, sobre conjuntos de hasta 200M de registros.',False)],
   [('Reglas de limpieza y calidad de datos en ',False),('TIBCO EBX',True),(': validaciones, restricciones y correcciones sobre registros maestros dentro de la plataforma.',False)],
   [('Extensiones Java sobre EBX: triggers, constraints, validaciones personalizadas, servicios REST, scripts de workflow e importadores — creados desde cero o evolucionados a partir de módulos existentes.',False)],
   [('Integración de EBX con sistemas externos mediante ',False),('SnapLogic iPaaS',True),(' y servicios REST nativos de EBX: Oracle, PostgreSQL, SQL Server y AWS.',False)],
   [('Telecomunicaciones: diseño de reglas de limpieza y ',False),('Match & Merge',True),(' para datos de clientes de dos países centroamericanos, tras investigar la numeración telefónica y los estándares de direcciones de cada país; entrega de datos limpios y deduplicados a la operadora.',False)],
   [('Monitoreo y reporte de calidad de datos con ',False),('Power BI',True),(' y ',False),('SQL',True),(' para dar seguimiento a KPIs de gobierno en proyectos de clientes.',False)],
   [('Cinco implementaciones MDM en ',False),('retail, manufactura, servicios financieros, farmacéutico y telecomunicaciones',True),(', trabajando con equipos de negocio y TI para traducir requerimientos a configuraciones técnicas en EBX.',False)]],
  'r2': ('Becario – Consultor de Datos y MDM','Alldatum Business  ·  Ciudad de México','Dic 2021 – Ago 2023','1 año 9 meses'),
  'r2b': [[('Aprendizaje y configuración inicial de la plataforma TIBCO EBX bajo supervisión de consultores senior.',False)],
          [('Apoyo en la identificación de fuentes de datos y su mapeo a modelos MDM durante proyectos de implementación.',False)]],
  'r3': ('Instructor de programación','PROTECO  ·  Facultad de Ingeniería, UNAM','Feb 2021 – Feb 2022','1 año 1 mes'),
  'r3b': [[('Impartí 10+ cursos y talleres en línea (POO, desarrollo web, Python, Linux) a grupos de 20–30 estudiantes de Ingeniería en Computación, incluyendo la inducción de nuevos integrantes del programa.',False)],
          [('Creación de material didáctico y tutorías individuales enfocadas en resolución de problemas técnicos, además de talleres prácticos como armar un CV profesional.',False)]],
  'h_edu': 'Educación',
  'e1': ('Ingeniero en Computación','Facultad de Ingeniería · Universidad Nacional Autónoma de México (UNAM)','2017 – 2024','Título y cédula 2024'),
  'e2': ('Técnico en Desarrollo de Software','Colegio de Ciencias y Humanidades, Azcapotzalco · UNAM','Ago 2014 – Mar 2017','Egresado'),
  'h_cert': 'Certificaciones',
  'certs': [[('Google AI',True),(' — Google · Ago 2026 · ID de credencial LY3LGZVG593W',False)],
            [('Diplomado: Desarrollo de Habilidades Directivas (240 h)',True),(' — Facultad de Ingeniería, UNAM · Mar 2024 · Folio 4283',False)],
            [('SnapLogic Integrator Certification',True),(' — SnapLogic · Ago 2023 · verify.skilljar.com/c/ewn4n2hoqxh5',False)],
            [('Google Data Analytics Professional Certificate',True),(' — Google / Coursera · Mar 2023 · coursera.org/share/ef3d02dd5eafc5509d940ecdb1be61b2',False)],
            [('Microsoft Certified: Azure Fundamentals (AZ-900)',True),(' — Microsoft vía Certiport · Jun 2020',False)]],
  'udemy': [('Formación continua (Udemy, 2025): ',True),('Curso Práctico SQL Server: De Principiante a Experto · SQL Consultas en SQL Server · Administrador Linux – Curso completo desde cero (LPIC-1) · Ciberseguridad – CompTIA Security+ Módulo 1 · Ciberseguridad: Protege tu información de cibercriminales. Formación oficial en la plataforma TIBCO EBX.',False)],
  'h_lang': 'Idiomas',
  'langs': 'Español — Nativo     |     Inglés — B2 (competencia técnica)',
 },
}[LANG]

# ── Documento ─────────────────────────────────────────────────────
BLUE=RGBColor(0x1F,0x5F,0xA8); GRAY=RGBColor(0x55,0x55,0x55); DARK=RGBColor(0x22,0x22,0x22)
doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(0.6); sec.top_margin=sec.bottom_margin=Inches(0.5)
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.space_before=Pt(0)

def para(text='',size=10,bold=False,color=DARK,align=None,after=0,before=0):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if align: p.alignment=align
    if text:
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=color
    return p
def runs(p,parts):
    for t,b in parts:
        r=p.add_run(t); r.bold=b; r.font.size=Pt(10); r.font.color.rgb=DARK
def bottom_border(p):
    pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'8'); b.set(qn('w:space'),'1'); b.set(qn('w:color'),'1F5FA8'); pbdr.append(b); pPr.append(pbdr)
def heading(text):
    p=para(text.upper(),size=10.5,bold=True,color=BLUE,before=9,after=3); bottom_border(p)
def bullet(parts):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1.5); p.paragraph_format.left_indent=Inches(0.22)
    runs(p,parts)
def role(title,org,dates,duration):
    # Tabla de 2x2: título/organización a la izquierda, fechas/duración a la derecha
    t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    for row in t.rows: row.cells[0].width=Inches(5.6); row.cells[1].width=Inches(1.7)
    c=t.cell(0,0).paragraphs[0]; r=c.add_run(title); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=DARK
    c=t.cell(1,0).paragraphs[0]; r=c.add_run(org); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=BLUE
    c=t.cell(0,1).paragraphs[0]; c.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=c.add_run(dates); r.font.size=Pt(9.5); r.font.color.rgb=GRAY
    c=t.cell(1,1).paragraphs[0]; c.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=c.add_run(duration); r.font.size=Pt(9.5); r.bold=True; r.font.color.rgb=BLUE
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

para('MARIO HUARTE NOLASCO',size=20,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
para(T['role'],size=12,bold=True,color=BLUE,align=WD_ALIGN_PARAGRAPH.CENTER,after=3)
para(T['contact1'],size=9,color=GRAY,align=WD_ALIGN_PARAGRAPH.CENTER,after=1)
para(T['contact2'],size=9,color=GRAY,align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
runs(para(before=2,after=2),T['summary'])
heading(T['h_comp']); para(T['comp'],size=9.5,after=2)
heading(T['h_exp'])
for key in ('r1','r2','r3'):
    role(*T[key])
    for b in T[key+'b']: bullet(b)
heading(T['h_edu']); role(*T['e1']); role(*T['e2'])
heading(T['h_cert'])
for b in T['certs']: bullet(b)
runs(para(before=3,after=1),T['udemy'])
heading(T['h_lang']); para(T['langs'])
doc.save(OUT); print('docx ok', OUT)
